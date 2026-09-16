#!/usr/bin/env python3
"""Check the built deck.

Two layers:

1. OOXML schema validation, delegated to the pptx skill's validator at
   /mnt/skills/public/pptx/scripts/office/validate.py when that tree is
   present. It is not vendored here because it needs its sibling modules
   and the XSD schema set. If it is missing the step is skipped with a note.

2. Deck-level checks that are specific to this deck and always run:
   every slide has a footer, its own page number and a non-empty notes
   pane; the notes match <build>/content.json exactly under the same
   "label — URL" transform build_deck.js applies; no emoji anywhere.

Usage: python3 validate.py [path/to/deck.pptx]
"""
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PLANS = os.environ.get("E080_PLANS", os.path.dirname(HERE))
BUILD = os.environ.get("E080_BUILD", os.path.join(HERE, "build"))
PPTX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    PLANS, "E080_hourglass_cone_deck.pptx")

FOOTER = "E-080 · the cut-off mirrored double cone"
SKILL = "/mnt/skills/public/pptx/scripts/office/validate.py"
EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿️]")

fail = []

if os.path.exists(SKILL):
    r = subprocess.run([sys.executable, SKILL, PPTX],
                       capture_output=True, text=True)
    tail = (r.stdout or r.stderr).strip().splitlines()[-1:] or [""]
    print("schema:", tail[0])
    if r.returncode != 0:
        fail.append("OOXML schema validation failed")
else:
    print(f"schema: skipped ({SKILL} not present)")

from pptx import Presentation                                  # noqa: E402

prs = Presentation(PPTX)
content = json.load(open(os.path.join(BUILD, "content.json"), encoding="utf8"))
n_expected = len(content["notes"])

if len(prs.slides) != n_expected:
    fail.append(f"{len(prs.slides)} slides, spec has {n_expected}")

for i, s in enumerate(prs.slides, start=1):
    texts = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame]
    notes = s.notes_slide.notes_text_frame.text.strip()
    want = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 — \2",
                  content["notes"].get(str(i), ""))
    if FOOTER not in texts:
        fail.append(f"slide {i}: no footer")
    if str(i) not in texts:
        fail.append(f"slide {i}: no page number")
    if not notes:
        fail.append(f"slide {i}: empty notes pane")
    elif notes != want:
        fail.append(f"slide {i}: notes differ from the spec")
    if EMOJI.search("".join(texts) + notes):
        fail.append(f"slide {i}: emoji in the text")

if fail:
    print("FAILED:")
    for f in fail:
        print("  -", f)
    sys.exit(1)
print(f"deck: {len(prs.slides)} slides, notes verbatim, footers and numbers OK")

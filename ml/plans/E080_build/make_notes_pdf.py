#!/usr/bin/env python3
"""E-080 deck with speaker notes: 40 pages, slide then notes, 16:9, dark."""
import glob
import os
import re

from pptx import Presentation
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
PLANS = os.environ.get("E080_PLANS", os.path.dirname(HERE))
BUILD = os.environ.get("E080_BUILD", os.path.join(HERE, "build"))
PPTX = f"{PLANS}/E080_hourglass_cone_deck.pptx"
OUT = f"{PLANS}/E080_hourglass_cone_deck_with_notes.pdf"

W, H = 960.0, 540.0                       # 13.333 x 7.5 in, in points
BG = HexColor("#0D1117")
CARD = HexColor("#161B22")
LINE = HexColor("#30363D")
TXT = HexColor("#E6EDF3")
MUT = HexColor("#7D8590")
FOOT = HexColor("#4A5460")
BLUE = HexColor("#4493F8")

prs = Presentation(PPTX)
slides = []
for s in prs.slides:
    shapes = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = shapes[1].text_frame.text.strip()
    notes = s.notes_slide.notes_text_frame.text.strip()
    assert title and notes
    slides.append((title, notes))
assert len(slides) == 20, len(slides)

pngs = sorted(glob.glob(f"{BUILD}/slidepage-*.png"))
assert len(pngs) == 20, pngs

c = canvas.Canvas(OUT, pagesize=(W, H))
c.setTitle("The cut-off mirrored double cone — slides with speaker notes")

for i, ((title, notes), png) in enumerate(zip(slides, pngs), start=1):
    # ---- slide page
    c.setFillColor(BG)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.drawImage(png, 0, 0, width=W, height=H)
    c.showPage()

    # ---- notes page
    c.setFillColor(BG)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(BLUE)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(43, H - 40, f"SPEAKER NOTES  ·  SLIDE {i} OF 20")
    cx, cy, cw = 43.0, 48.0, W - 86.0
    # the title wraps to a second line when it is too long for one
    tstyle = ParagraphStyle("t", fontName="Times-Bold", fontSize=23,
                            leading=27, textColor=TXT)
    tpara = Paragraph(title.replace("&", "&amp;").replace("<", "&lt;"), tstyle)
    tw, th = tpara.wrap(cw, 200)
    tpara.drawOn(c, cx, H - 52 - th)
    ch = (H - 52 - th - 14) - cy
    c.setFillColor(CARD)
    c.setStrokeColor(LINE)
    c.setLineWidth(1)
    c.roundRect(cx, cy, cw, ch, 6, stroke=1, fill=1)

    pad = 26.0
    size = 16.0
    body = notes.replace("&", "&amp;").replace("<", "&lt;")
    # markdown links become clickable, accent blue, underlined
    body = re.sub(r"\[([^\]]+)\]\(([^)]+)\)",
                  r'<a href="\2"><font color="#4493F8"><u>\1</u></font></a>', body)
    while size >= 10.0:
        st = ParagraphStyle("n", fontName="Helvetica", fontSize=size,
                            leading=size * 1.42, textColor=TXT)
        para = Paragraph(body, st)
        pw, ph = para.wrap(cw - 2 * pad, ch - 2 * pad)
        if ph <= ch - 2 * pad:
            break
        size -= 0.5
    para.drawOn(c, cx + pad, cy + (ch - ph) / 2.0)

    c.setFillColor(FOOT)
    c.setFont("Helvetica", 10)
    c.drawString(43, 26, "E-080 · the cut-off mirrored double cone — speaker notes")
    c.drawRightString(W - 43, 26, f"{i}")
    c.showPage()

c.save()
print("wrote", OUT)

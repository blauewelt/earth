"""The synthetic archive of family 1.3's `mplastic` and `mobjects` (E-092).

Every source in its REAL on-disk layout, mirrored by URL exactly as the
adapters fetch it (`mplastic.mirror_name`): the two figshare items (API
listings, the 2018 .xls files with float-noise event ids and the Megaplastics
sheet's latitude/longitude SWAPPED, the 2024 .xlsx with its description), the
NCEI Hub CSV with the real 37-column header, the NCEI landing pages, archive
listings and accession spreadsheets, and the EMODnet zip with the real
113-column ODV header.

Hostile on purpose, and every case is a named expectation in the tests:
a measured zero in every source that has one; NCEI "Law et al. 2014" rows
that must be dropped BY NAME and "Egger et al. 2020" rows dropped only BY
MATCH (+-1 day, 0.06 deg) or kept; an NCEI study matched
1:1 to its own accession (method a) with two tows at one station on one day;
a study back-converted by a measured depth (method b, 25 matched rows) and one
left per m3 (method c); a pump sample; beach and nurdle rows that are not
water; an NCEI row and an EMODnet sample that are copies of earlier sources;
EMODnet lines that are not plastic, ambiguous, VOR-coded or date-only; a 2024
row with no value. The truth is computed HERE from the generator's own
parameters, never by calling the adapter.

Imported only by the smoke (it needs `xlwt` to write .xls, which the build
does not).
"""
import csv
import io
import json
import math
import os
import zipfile

import numpy as np

import build_family10_stores as f10b
from family1.adapters import _lebreton2018 as lb
from family1.adapters import mplastic as mp

NAN = float("nan")
KN = 0.514444


def L(x):
    """The stored form of a concentration: log10(1 + x), NaN for None."""
    return NAN if x is None else math.log10(1.0 + x)


def _w(root, url, data):
    p = os.path.join(root, "mplastic", mp.mirror_name(url))
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "wb") as fh:
        fh.write(data if isinstance(data, bytes) else data.encode("utf-8"))
    return p


def _t(y, m, d, hour=None):
    base = int((np.datetime64(f"{y:04d}-{m:02d}-{d:02d}") -
                np.datetime64("1982-01-01")).astype(int)) * 86400
    if hour is None:
        return base + 43200
    return base + min(int(round(hour * 60)), 1439) * 60


def _t_mid_month(y, m):
    import calendar
    start = _t(y, m, 1) - 43200
    return start + calendar.monthrange(y, m)[1] * 86400 // 2


def _xls(sheets):
    """{sheet: (header, rows)} -> .xls bytes (xlwt)."""
    import xlwt
    wb = xlwt.Workbook()
    for name, (hdr, rows) in sheets.items():
        ws = wb.add_sheet(name)
        for j, h in enumerate(hdr):
            ws.write(0, j, h)
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r):
                if v is None:
                    continue
                ws.write(i, j, v)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _xlsx(sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (hdr, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(hdr))
        for r in rows:
            ws.append(list(r))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ============================================================ Lebreton 2018 ==
N_MANTA, N_MEGA, N_AER = 501, 151, 31
STATION_HDR = ['#', 'Sampling event ID', 'Sampling type', 'Platform name',
               'Start day (UTC)', 'Start month (UTC)', 'Start year (UTC)',
               'Start UTC time (hours) ', 'End UTC time (hours)',
               'Sampling duration (hours)', 'Start latitude (degrees)',
               'Start longitude (degrees)', 'End latitude (degrees)',
               'End longitude (degrees)', 'Sampling distance (km)',
               'Sampling width (km)', 'Sampling area (km2)',
               'Sampling depth (km)', 'Wind speed (Knots)',
               'Sea state (Beaufort)', 'Vessel speed (knots)', 'Observations']
CONC_HDR = ['#', 'Sampling event ID', 'Longitude (degrees)',
            'Latitude (degrees)', 'Sampling type',
            'Midpoint Estimate (# km-2)', 'Midpoint Estimate (g km-2)',
            'Lower Estimate (# km-2)', 'Lower Estimate  (g km-2)',
            'Higher Estimate (# km-2)', 'Higher Estimate (g km-2)',
            'Raw Data (# km-2)', 'Raw data (g km-2)']
MOSAIC_HDR = ['#', 'Sampling event ID',
              'Bin ID (at The Ocean Cleanup database)',
              'Start longitude debris (degrees)',
              'Start latitude debris (degrees)',
              'End longitude debris (degrees)',
              'End latitude debris (degrees)',
              'Flight line ID (at The Ocean Cleanup database)',
              'Image ID (at The Ocean Cleanup database)', 'Object type',
              'Color', 'Length (m)', 'Width (m)', '< 50 cm?',
              'Object distinctiveness ', 'Container length (m)',
              'Container width (m)', 'Topview area (m^2)',
              'Lower weight (kg)', 'Mid weight (kg)', 'Upper weight (kg)',
              'Observations']
HIST_HDR = ['#', 'Data Origin', 'Data Reference', 'Sampling Year',
            'Sampling Month', 'Longitude (degrees)', 'Latitude (degrees)',
            'Sampling type', 'Mesh Size (microns)', 'Net tow Depth (in m)',
            'Microplastic Numerical Concentration (#/km²)',
            'Microplastic Mass Concentration (g/km²)', 'Inside GPGP?']
OBJ_TYPES = ("Bundled net", "Loose net", "Rope", "Buoy", "Container",
             "Life ring", "Unknown")
ZERO_MANTA = 3                 # the manta tow whose net came up empty
N_THIS = 501                   # 'This study' rows = the manta events again
N_LAW14 = 2529
N_ERIK = 180                   # Eriksen 2014, gear blank as in the real file
N_ERIK_MATCH = 25              # of which NCEI holds these (method b)
N_OTHER = 3532 - N_THIS - N_LAW14 - N_ERIK


def events():
    """The 683 synthetic events (StationInfo order) with their designed
    values. Positions are 0.03 deg apart, so nothing matches by accident."""
    out = []
    for k in range(N_MANTA + N_MEGA + N_AER):
        if k < N_MANTA:
            typ, y, mo = "manta", 2015, 7 + (k % 2)
            d = 1 + (k // 2) % 28
            hour = (k % 20) + 0.25
        elif k < N_MANTA + N_MEGA:
            typ, y, mo = "mega", 2015, 8
            d = 1 + (k % 28)
            hour = 3.5 + (k % 10)
        else:
            typ, y, mo = "aerial", 2016, 10
            d = 2 if k % 2 else 6
            hour = None
        la0 = 20.0 + 0.03 * k
        lo0 = -150.0 + 0.5 * (k % 7)
        e = {"k": k, "type": typ, "y": y, "mo": mo, "d": d, "hour": hour,
             "lat0": la0, "lon0": lo0, "lat1": la0 + 0.004,
             "lon1": lo0 + 0.004, "platform": f"Vessel {k % 3}",
             # stored with float noise; the sheets carry the rounded id
             "sid": k * 0.1 + (2e-16 if k % 3 else 0.0),
             "cid": round(k * 0.1, 1)}
        e["lat"] = round((e["lat0"] + e["lat1"]) / 2, 4)
        e["lon"] = round((e["lon0"] + e["lon1"]) / 2, 4)
        if typ == "manta":
            z = 0.0 if k == ZERO_MANTA else 1.0
            e["micro"] = dict(raw_n=z * 1000.0 * (k + 1), mid_n=z * 1500.0 *
                              (k + 1), raw_g=z * 2.0 * (k + 1),
                              mid_g=z * 3.0 * (k + 1))
            e["meso"] = dict(raw_n=z * 10.0 * (k + 1), mid_n=z * 12.0 *
                             (k + 1), raw_g=z * 1.0 * (k + 1),
                             mid_g=z * 1.5 * (k + 1))
            e["depth_km"] = 0.00015
        elif typ == "mega":
            e["macro"] = dict(raw_n=50.0 + k, mid_n=50.0 + k,
                              raw_g=200.0 + k, mid_g=200.0 + k)
            e["depth_km"] = 0.001
        else:
            e["mega"] = dict(raw_n=2.0 + 0.1 * (k - 652), mid_n=2.0 + 0.1 *
                             (k - 652), raw_g=None, mid_g=1000.0 + k)
            e["depth_km"] = None
            e["area_km2"] = 4.0
        out.append(e)
    return out


def historical(evs):
    """3,532 rows: 'This study' (= the manta events), Law 2014, Eriksen 2014,
    and 'Wong et al. 1974' filler. Returns (rows, the Eriksen rows NCEI will
    hold)."""
    rows = []
    manta = [e for e in evs if e["type"] == "manta"]
    for i, e in enumerate(manta):
        n = e["micro"]["raw_n"] + e["meso"]["raw_n"]
        rows.append({"origin": "MEGA expedition", "ref": "This study",
                     "y": e["y"], "mo": e["mo"], "lon": e["lon"],
                     "lat": e["lat"], "gear": "manta", "mesh": 500.0,
                     "depth": 0.15, "n": n, "g": 1.0, "inside": 1.0})
    for j in range(N_LAW14):
        rows.append({"origin": f"SEA cruise {j % 40}",
                     "ref": "Law et al. 2014", "y": 2001 + j % 12,
                     "mo": 1 + j % 12, "lon": 170.0 - 0.02 * (j % 500),
                     "lat": -10.0 + 0.02 * (j // 500) * 50 + 0.0001 * j,
                     "gear": "neuston", "mesh": 335.0, "depth": "nan",
                     "n": 0.0 if j % 4 == 0 else 100.0 * j, "g": 0.5 * j,
                     "inside": 0.0})
    erik = []
    for j in range(N_ERIK):
        r = {"origin": "5 Gyres", "ref": "Eriksen et al.  2014", "y": 2013,
             "mo": 1 + j % 12, "lon": 60.0 + 0.05 * j, "lat": -30.0 + 0.0 * j,
             "gear": "", "mesh": 333.0, "depth": "nan",
             "n": 4000.0 + 10.0 * j, "g": 1.0 + j, "inside": 0.0}
        rows.append(r)
        if j < N_ERIK_MATCH:
            erik.append(r)
    for j in range(N_OTHER):
        rows.append({"origin": f"Origin {j % 9}", "ref": "Wong et al. 1974",
                     "y": 1972 + j % 3, "mo": 4 + j % 6,
                     "lon": -170.0 + 0.05 * j, "lat": 45.0,
                     "gear": "neuston" if j % 2 else "manta",
                     "mesh": 333.0, "depth": 0.25 if j % 2 else "nan",
                     # 77 of the real rows give a mass and no count
                     "n": None if j % 10 == 0 else 20.0 * j,
                     "g": 0.0 if j == 1 else 0.1 * j, "inside": 0.0})
    return rows, erik


def objects(evs):
    aer = [e for e in evs if e["type"] == "aerial"]
    out = []
    for o in range(lb.EXPECT_OBJECTS):
        e = aer[o % len(aer)]
        lt50 = "yes" if o % 3 == 0 else "no"
        out.append({"e": e, "lon0": e["lon"] + 0.001 * (o % 5),
                    "lat0": e["lat"] + 0.0005,
                    "lon1": e["lon"] + 0.001 * (o % 5) + 0.0002,
                    "lat1": e["lat"] + 0.0007,
                    "type": OBJ_TYPES[o % len(OBJ_TYPES)],
                    "length": 0.3 + 0.01 * (o % 50),
                    "width": 0.1 + 0.005 * (o % 20), "lt50": lt50,
                    "area": None if o == 7 else 0.02 + 0.001 * (o % 40),
                    "kg": (0.1 * (o % 9), 0.2 * (o % 9), 0.3 * (o % 9))})
    return out


def write_lebreton2018(root, evs, hist, objs, corrupt_swap=False):
    st_rows = []
    for e in evs:
        typ = {"manta": "Manta trawl", "mega": "Mega trawl",
               "aerial": "RGB mosaic"}[e["type"]]
        st_rows.append([e["k"] + 1, e["sid"], typ, e["platform"], e["d"],
                        e["mo"], e["y"],
                        "" if e["hour"] is None else e["hour"], "",
                        1.0, e["lat0"], e["lon0"], e["lat1"], e["lon1"],
                        7.3, 0.0009, e.get("area_km2", 0.0066),
                        "" if e["depth_km"] is None else e["depth_km"],
                        10.0, 3.0, 3.0, ""])
    sheets = {}
    for key, sheet, _n, stype in lb.SHEETS:
        rows = []
        for e in evs:
            if key not in e:
                continue
            c = e[key]
            lo_, la_ = e["lon"], e["lat"]
            if key == "mega" and not (corrupt_swap and e["k"] == 660):
                lo_, la_ = la_, lo_          # the published file's trap
            g_raw = "nan" if c["raw_g"] is None else c["raw_g"]
            rows.append([len(rows) + 1, e["cid"], lo_, la_, stype,
                         c["mid_n"], c["mid_g"], c["raw_n"], c["mid_g"],
                         c["mid_n"], c["mid_g"], c["raw_n"], g_raw])
        sheets[sheet] = (CONC_HDR, rows)
    mo_rows = []
    for i, o in enumerate(objs):
        mo_rows.append([i + 1, o["e"]["cid"], 1.1, o["lon0"], o["lat0"],
                        o["lon1"], o["lat1"], 185604.0, f"IMG{i}",
                        o["type"], "White", o["length"], o["width"],
                        o["lt50"], "Good", "", "",
                        "" if o["area"] is None else o["area"],
                        o["kg"][0], o["kg"][1], o["kg"][2], ""])
    h_rows = []
    for i, h in enumerate(hist):
        h_rows.append([i + 1, h["origin"], h["ref"], h["y"], h["mo"],
                       h["lon"], h["lat"], h["gear"], h["mesh"], h["depth"],
                       "nan" if h["n"] is None else h["n"], h["g"],
                       h["inside"]])
    files = {
        "Lebreton2018_Concentration.xls": (10441962, _xls(sheets)),
        "Lebreton2018_SamplingInformation.xls": (10441968, _xls({
            "StationInfo": (STATION_HDR, st_rows),
            "TrawlDebrisInfo": (["#"], []),
            "MosaicDebrisInfo": (MOSAIC_HDR, mo_rows)})),
        "Lebreton2018_HistoricalDataset.xls": (10441965, _xls({
            "HistoricalDataset_3.2": (HIST_HDR, h_rows)})),
        "GPGP_contours.zip": (44299727, b"decoy"),
    }
    listing = []
    for name, (fid, data) in files.items():
        url = f"https://ndownloader.figshare.com/files/{fid}"
        _w(root, url, data)
        listing.append({"id": fid, "name": name, "size": len(data),
                        "download_url": url})
    _w(root, lb.API_FILES, json.dumps(listing))


# ============================================================ Lebreton 2024 ==
def _desc():
    names = {1: "small microplastics", 2: "large microplastics",
             3: "small mesoplastics", 4: "large mesoplastics",
             5: "small macroplastics", 6: "large macroplastics",
             7: "megaplastics"}
    items = []
    for k, (u, kind) in (("CD", ("#/km", "numerical")),
                         ("WD", ("g/km", "mass"))):
        for n, (lo, hi) in mp.L2024_CLASS_MM.items():
            rng = f"> {lo:g} mm" if hi is None else f"{lo:g} mm - {hi:g} mm"
            items.append(f"<li><b>{k}{n}</b>: {kind} concentration of "
                         f"{names[n]} ({rng}) in {u}<sup>2</sup></li>")
    items.append("<li><b>WD567</b>: mass concentration of macroplastics and "
                 "megaplastics (> 50 mm) in g/km<sup>2</sup></li>")
    return ("<p>Surface concentration of floating plastic pollution in the "
            "North Pacific Ocean measured from 2015 to 2022 (manta trawl, mega "
            "trawl, aerial survey, and cleanup system)</p><ul>" +
            "".join(items) + "</ul>")


def lebreton2024_book(evs):
    """The four sheets and the truth of the rows kept from them."""
    base = ["Expedition", "Lat", "Lon", "Day", "Month", "Year", "In NPGP",
            "In Hotspots"]
    manta_h = base + [f"CD{i}" for i in range(1, 7)] + \
        [f"WD{i}" for i in range(1, 7)]
    mega_h = base + ["CD5", "CD6", "CD7", "WD5", "WD6", "WD7"]
    aer_h = base + ["CD7", "WD7"]
    clean_h = base + ["WD567"]
    truth = []
    manta = [e for e in evs if e["type"] == "manta"][:120]
    mr = []
    for e in manta:
        mid = e["micro"]["mid_n"]
        mr.append(["'MEGA Expedition'", round(e["lat"], 2), round(e["lon"], 2),
                   e["d"], e["mo"], e["y"], 0, 0, 0.6 * mid, 0.4 * mid,
                   e["meso"]["mid_n"], 0, 0, 0] + [0] * 6)
    for j in range(10):
        y, mo, d = 2019, 6, 1 + j
        cd = None if j == 9 else 100.0 * (j + 1)
        row = ["'North Pacific Mission 3 (NPM3)'", 35.0 + 0.1 * j, -140.0,
               d, mo, y, 1, 0, cd, 50.0, 5.0, 1.0, 0.0, 0.0] + [0.1] * 6
        mr.append(row)
        if cd is not None:
            truth.append(("lebreton2024", (0, len(mr) - 1),
                          dict(t=_t(y, mo, d), lat=35.0 + 0.1 * j,
                               lon=-140.0, qc=1 | mp.QC_NO_TIME,
                               plat="lebreton2024|North Pacific Mission 3 "
                                    "(NPM3)|",
                               v={"log_n_km2_wc": L(cd + 56.0),
                                  "sample_depth_m": 0.0})))
    megr = []
    for e in [e for e in evs if e["type"] == "mega"][:20]:
        raw = e["macro"]["raw_n"]
        megr.append(["'MEGA Expedition'", e["lat"], e["lon"], e["d"], e["mo"],
                     e["y"], 0, 0, 0.25 * raw, 0.75 * raw, 1.0, 1, 1, 1])
    for j in range(2):
        y, mo, d = 2022, 7, 20 + j
        megr.append(["'North Pacific Mission 5 (NPM5)'", 32.0, -133.0 + j, d,
                     mo, y, 1, 1, 600.0, 170.0, 30.0 * j, 1300.0, 8600.0,
                     5800.0])
        truth.append(("lebreton2024", (1, len(megr) - 1), dict(
            t=_t(y, mo, d), lat=32.0, lon=-133.0 + j, qc=3 | mp.QC_NO_TIME,
            plat="lebreton2024|North Pacific Mission 5 (NPM5)|",
            v={"log_n_km2": L(770.0 + 30.0 * j),
               "log_g_km2": L(1300.0 + 8600.0 + 5800.0),
               "log_n_km2_wc": L(770.0 + 30.0 * j),
               "log_n_km2_macro": L(770.0),
               "log_n_km2_mega": L(30.0 * j), "sample_depth_m": 0.0})))
    aerr = []
    for e in [e for e in evs if e["type"] == "aerial"]:
        aerr.append(["AERIAL Expedition'", e["lat"], e["lon"], e["d"],
                     e["mo"], e["y"], 1, 1, round(e["mega"]["raw_n"], 2),
                     e["mega"]["mid_g"]])
    for j in range(3):
        y, mo, d = 2021, 8, 5 + j
        cd7 = 0.0 if j == 0 else 3.5 * j
        aerr.append(["S002-1'", 34.9, -145.0 + j, d, mo, y, 1, 1, cd7, 900.0])
        truth.append(("lebreton2024", (2, len(aerr) - 1), dict(
            t=_t(y, mo, d), lat=34.9, lon=-145.0 + j, qc=6 | mp.QC_NO_TIME,
            plat="lebreton2024|S002-1|",
            v={"log_n_km2": L(cd7), "log_n_km2_mega": L(cd7),
               "sample_depth_m": 0.0})))
    clr = []
    for j in range(2):
        y, mo, d = 2022, 12, 1 + j
        clr.append(["S002'", 33.0, -138.0 + j, d, mo, y, 1, 1, 47000.0 + j])
        truth.append(("lebreton2024", (3, len(clr) - 1), dict(
            t=_t(y, mo, d), lat=33.0, lon=-138.0 + j, qc=8 | mp.QC_NO_TIME,
            plat="lebreton2024|S002|",
            v={"log_g_km2": L(47000.0 + j), "sample_depth_m": 0.0})))
    book = {"Manta trawl": (manta_h, mr), "Mega trawl": (mega_h, megr),
            "Aerial survey": (aer_h, aerr), "Cleanup system": (clean_h, clr)}
    return book, truth


def write_lebreton2024(root, evs):
    book, truth = lebreton2024_book(evs)
    data = _xlsx(book)
    url = "https://ndownloader.figshare.com/files/48791128"
    _w(root, url, data)
    _w(root, mp.FIG2024_ARTICLE, json.dumps({
        "id": 25674120, "title": "Supplementary data", "description": _desc(),
        "license": {"value": 1, "name": "CC BY 4.0",
                    "url": "https://creativecommons.org/licenses/by/4.0/"}}))
    _w(root, mp.FIG2024_FILES, json.dumps([
        {"id": 48791128, "name": mp.FIG2024_NAME, "size": len(data),
         "download_url": url}]))
    return truth


# ===================================================================== NCEI ==
NCEI_HDR = ['Unique ID', 'Sample Date', 'Latitude (degree)',
            'Longitude (degree)', 'Ocean', 'Region', 'Subregion', 'Country',
            'State', 'Beach Location', 'Marine Setting',
            'Ocean Bottom Depth (m)', 'Water Sample Depth (m)',
            'Sediment Sample Depth (m)', 'Sampling Method', 'Mesh Size (mm)',
            'Transect Number', 'Sampling Point on Beach', 'Volunteers Number',
            'Collecting Time (min)', 'Standardized Nurdle  Amount',
            'Microplastics Measurement', 'Unit', 'Concentration Class Range',
            'Concentration Class', 'Short Reference', 'Long Reference', 'DOI',
            'Organization', 'Keywords', 'NCEI Accession Number',
            'NCEI Accession Link', 'Study Type', 'ObjectId', 'Symbology', 'x',
            'y']
ACC_ARC = {"0211007": "arc0157", "0275968": "arc0201", "0314178": "arc0240",
           "0276264": "arc0202"}


def _ncei_row(uid, y, m, d, la, lo, setting, method, mesh_mm, val, unit,
              ref, org, acc, depth="0"):
    r = {h: "" for h in NCEI_HDR}
    r.update({"Unique ID": uid, "Sample Date": f"{m}/{d}/{y} 5:00:00 AM",
              "Latitude (degree)": f"{la}", "Longitude (degree)": f"{lo}",
              "Marine Setting": setting, "Water Sample Depth (m)": depth,
              "Sampling Method": method, "Mesh Size (mm)": f"{mesh_mm}",
              "Microplastics Measurement": f"{val}", "Unit": unit,
              "Short Reference": ref, "Organization": org,
              "NCEI Accession Number": acc})
    return r


def write_ncei(root, evs, erik):
    """The Hub CSV, three accessions, and the truth of the NCEI rows kept."""
    rows, truth = [], []

    def T(r, gear, v, depth_used=None):
        y, m, d = (int(x) for x in (r["Sample Date"].split(" ")[0]
                                    .split("/")[2], r["Sample Date"]
                                    .split("/")[0], r["Sample Date"]
                                    .split("/")[1]))
        study = r["Short Reference"]
        mesh = float(r["Mesh Size (mm)"]) * 1000
        sd = r["Water Sample Depth (m)"]
        vv = dict(v)
        vv.update(mesh_um=mesh, sample_depth_m=float(sd) if sd else None)
        if depth_used is not None:
            vv["net_depth_m"] = depth_used
        truth.append(("ncei", (study, r["Sampling Method"],
                               len(rows) - 1), dict(
            t=_t(y, m, d), lat=float(r["Latitude (degree)"]),
            lon=float(r["Longitude (degree)"]), qc=gear | mp.QC_NO_TIME,
            plat=f"ncei|{study}|{r['Organization']}", v=vv)))
    # --- Law et al. 2010 (method a): 6 tows, two at one station on one day
    law10 = [(1991, 10, 18, 35.88, -58.98, 17819.0),
             (1991, 10, 18, 35.88, -58.98, 18359.0),
             (1992, 5, 7, 32.1, -64.95, 0.0),
             (1993, 5, 7, 26.3, -78.37, 540.0),
             (1994, 6, 1, 30.0, -70.0, 1200.0),
             (1995, 6, 2, 31.0, -71.0, 2400.0)]
    acc_rows = []
    for i, (y, m, d, la, lo, km2) in enumerate(law10):
        acc_rows.append([np.datetime64(f"{y:04d}-{m:02d}-{d:02d}").astype(
            "datetime64[D]").item(), la, lo, km2])
    acc_rows.append([np.datetime64("1996-01-01").item(), 10.0, -10.0, 5.0])
    # NCEI lists the station pair in the OPPOSITE order: only the value
    # ratio can pair them correctly
    order = [1, 0, 2, 3, 4, 5]
    for i in order:
        y, m, d, la, lo, km2 = law10[i]
        r = _ncei_row(f"0211007-{i:07d}", y, m, d, la, lo, "Ocean water",
                      "Neuston net", 0.335, round(km2 / 0.25e6, 9),
                      "pieces/m3", "Law et al. 2010",
                      "Sea Education Association", "211007")
        rows.append(r)
        T(r, 2, {"log_n_km2": L(km2)}, depth_used=0.25)
    # --- Law et al. 2014: dropped by name (D9)
    for i in range(2):
        rows.append(_ncei_row(f"0211008-{i:07d}", 2005, 9, 1 + i, 15.0,
                              -160.0 + i, "Ocean water", "Neuston net",
                              0.335, 0.01, "pieces/m3", "Law et al. 2014",
                              "Sea Education Association", "211008"))
    # --- Egger et al. 2020, The Ocean Cleanup's own submission: dropped by
    #     MATCH, not by name — a copy of a manta event on its day, a copy a
    #     day later and 0.03 deg off the event's END (inside +-1 day and
    #     0.06 deg, outside the generic day + 0.01 deg rule), and a 2018 tow
    #     no Ocean Cleanup file holds, which is KEPT (per m3: no copy to
    #     measure a depth on)
    e0 = [e for e in evs if e["type"] == "manta"][5]
    egger = [(e0["y"], e0["mo"], e0["d"], e0["lat"], e0["lon"]),
             (e0["y"], e0["mo"], e0["d"] + 1, e0["lat1"] + 0.03,
              e0["lon1"] - 0.03),
             (2018, 9, 12, 36.3, -127.7)]
    for i, (y, m, d, la, lo) in enumerate(egger):
        r = _ncei_row(f"0276264-{i:07d}", y, m, d, la, lo, "Ocean water",
                      "Manta net", 0.5, 0.88, "pieces/m3",
                      "Egger et al. 2020",
                      "The Ocean Cleanup, Rotterdam, The Netherlands",
                      "276264")
        rows.append(r)
        if i == 2:
            T(r, 1, {"log_n_m3": L(0.88)})
    # --- Eriksen et al. 2014 (method b): 25 copies of the compiled rows give
    #     the depth 0.08 m; 5 more are back-converted with it and kept
    for i, h in enumerate(erik):
        rows.append(_ncei_row(f"0275968-{i:07d}", h["y"], h["mo"], 10,
                              h["lat"], h["lon"], "Ocean water", "Neuston net",
                              0.333, h["n"] / 0.08e6, "pieces/m3",
                              "Eriksen et al. 2014", "5 Gyres Institute",
                              "275968"))
    for i in range(5):
        v = 0.0 if i == 0 else 0.5 * i
        r = _ncei_row(f"0275968-{100 + i:07d}", 2012, 3, 1 + i, -40.0,
                      100.0 + i, "Ocean water", "Neuston net", 0.333, v,
                      "pieces/m3", "Eriksen et al. 2014", "5 Gyres Institute",
                      "275968")
        rows.append(r)
        T(r, 2, {"log_n_km2": L(v * 0.08e6)}, depth_used=0.08)
    # --- Doyle et al. 2011 (method c): per m3 only; a measured zero; one a
    #     copy of a manta event (dropped: lebreton2018 > ncei)
    for i in range(3):
        v = 0.0 if i == 1 else 0.02 * (i + 1)
        r = _ncei_row(f"0314178-{i:07d}", 2007, 8, 10 + i, 40.0 + i, -130.0,
                      "Ocean water", "Neuston net", 0.335, v, "pieces/m3",
                      "Doyle et al. 2011", "University of Washington",
                      "314178")
        rows.append(r)
        T(r, 2, {"log_n_m3": L(v)})
    e1 = [e for e in evs if e["type"] == "manta"][10]
    rows.append(_ncei_row("0314178-0000099", e1["y"], e1["mo"], e1["d"],
                          e1["lat"] + 0.005, e1["lon"] - 0.005, "Ocean water",
                          "Neuston net", 0.335, 0.3, "pieces/m3",
                          "Doyle et al. 2011", "University of Washington",
                          "314178"))
    # --- Lusher et al. 2014: a pump (volumetric), not crawled
    for i in range(3):
        v = 2.5 * (i + 1)
        r = _ncei_row(f"0314176-{i:07d}", 2013, 4, 1 + i, 55.0 + i, -10.0,
                      "Ocean water", "Pump", 0.25, v, "pieces/m3",
                      "Lusher et al. 2014", "GMIT", "314176", depth="3")
        rows.append(r)
        T(r, 4, {"log_n_m3": L(v)})
    # --- not water
    rows.append(_ncei_row("B-1", 2019, 1, 1, 20.0, 20.0, "Beach",
                          "Hand picking", 1, 3.0, "pieces/m3", "Beach study",
                          "X", "999999"))
    rows.append(_ncei_row("N-1", 2019, 1, 2, 27.0, -97.0, "Nurdle Patrol",
                          "Hand picking", 1, 3.0, "pieces/10 mins", "Nurdle",
                          "X", "999998"))
    bio = io.StringIO()
    w = csv.DictWriter(bio, fieldnames=NCEI_HDR, lineterminator="\r\n")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    _w(root, mp.NCEI_CSV, "﻿" + bio.getvalue())
    # --- the accessions, through landing page -> archive -> version -> files
    acc_files = {
        "0211007": {"Microplastic_Data_SEA_NorthAtlantic.xlsx": _xlsx({
            "SEA": (["Sample Date", "Lat(deg)", "Long(deg)", "Pieces/KM2"],
                    acc_rows)}),
            "SEA_PlasticData_Permission.pdf": b"%PDF decoy"},
        "0275968": {"MP_data_S2N.xlsx": _xlsx({
            "Data_Dictionary": (["Column Name", "Definition"],
                                [["Date", "date in ISO 8601"]]),
            "MP_data_S2N": (["Date", "Latitude", "Longitude",
                             "Microplastics measurement ", "Unit"],
                            [["2013-01-10", -30.0, 60.0, 0.05,
                              "pieces/m3"]])})},
        "0276264": {"MP_data_S2N.xlsx": _xlsx({
            "MP_data_S2N": (["Date", "Latitude", "Longitude",
                             "Microplastics measurement", "Unit"],
                            [["2018-09-12", 36.3, -127.7, 0.88,
                              "pieces/m3"]])})},
        "0314178": {"MP_data_S2N_Doyle_MER2011.xlsx": _xlsx({
            "Microplastics_Data": (["Date", "Latitude", "Longitude",
                                    "Microplastic Abundance (MPs/m3)"],
                                   [["2007-08-10", 40.0, -130.0, 0.02]])})},
    }
    for acc7, files in acc_files.items():
        arc = ACC_ARC[acc7]
        _w(root, mp.NCEI_LANDING.format(acc7=acc7),
           f'<html><a href="ftp://ftp-oceans.ncei.noaa.gov/nodc/archive/'
           f'{arc}/{acc7}/">ftp</a></html>')
        base = mp.NCEI_ARCHIVE.format(arc=arc, acc7=acc7)
        _w(root, base, '<a href="?C=N;O=D">Name</a><a href="/data/oceans/">'
                       'Parent Directory</a><a href="1.1/">1.1/</a>'
                       '<a href="1.2/">1.2/</a>')
        d0 = f"{base}1.2/data/0-data/"
        links = "".join(f'<a href="{n}">{n}</a>' for n in sorted(files))
        _w(root, d0, '<a href="/data/oceans/">Parent</a>' + links)
        for n, data in files.items():
            _w(root, d0 + n, data)
    return truth


# ================================================================== EMODnet ==
EMOD_HDR = ['Cruise', 'Station', 'Type', 'yyyy-mm-ddThh:mm:ss.sss',
            'Longitude [degrees_east]', 'Latitude [degrees_north]',
            'LOCAL_CDI_ID', 'EDMO_code', 'Bot. Depth [m]', 'Instrument Info',
            'Codes in Originator File', 'P35 Contributor Codes', 'References',
            'Comments', 'Data set name', 'Discipline', 'Category',
            'Variables measured', 'Data format', 'Data format version',
            'Data size', 'Data set creation date', 'Datum',
            'Measuring area type', 'Water depth [m]', 'Depth reference',
            'Minimum instrument depth [m]', 'Maximum instrument depth [m]',
            'Start date', 'Start time', 'End date', 'End time',
            'Vertical resolution', 'Vertical resolution unit',
            'Instrument / gear type', 'Track resolution',
            'Track resolution unit', 'Frequency', 'Frequency unit',
            'Platform type', 'Cruise name', 'Alternative cruise name',
            'Cruise start date', 'Station name', 'Alternative station name',
            'Station start date', 'Originator', 'Data Holding centre',
            'Project name', 'Project type', 'EDMED references',
            'CSR references', 'CSR platform name', 'CSR platform code',
            'Publication references', 'Data Quality Control Info',
            'Data Distributor', 'Database reference',
            'Access/ordering of data', 'Access restriction', 'CDI-record id',
            'CDI-record creation date', 'CDI-partner',
            'Location qualification', 'MinDepBelowSeaSurf [m]',
            'QV:SEADATANET', 'MaxDepBelowSeaSurf [m]', 'QV:SEADATANET',
            'End_Lat', 'QV:SEADATANET', 'End_Lon', 'QV:SEADATANET',
            'EventEndDateTime [yyyy-mm-ddThh:mm:ss]', 'QV:SEADATANET',
            'mouth_width [cm]', 'QV:SEADATANET', 'Mouth_area [cm2]',
            'QV:SEADATANET', 'Mesh_size [micrometres]', 'QV:SEADATANET',
            'Sampling_protocol', 'QV:SEADATANET', 'SampTrkLen [km]',
            'QV:SEADATANET', 'VolWBodySamp [m3]', 'QV:SEADATANET', 'sampID',
            'QV:SEADATANET', 'SubsampleID', 'QV:SEADATANET',
            'Microlitter_count_water', 'QV:SEADATANET',
            'Microlitter_count_water_VOR', 'QV:SEADATANET',
            '*Microlitter_weight_water [g]', 'QV:SEADATANET',
            'Microlitter_type_water', 'QV:SEADATANET',
            'Microlitter_size_class_water', 'QV:SEADATANET',
            'Microlitter_polymer_type_water', 'QV:SEADATANET',
            'Microlitter_colour_class_water', 'QV:SEADATANET',
            'Microlitter_transparency_water', 'QV:SEADATANET', 'WMOSeaState',
            'QV:SEADATANET', 'WindSpd [m/s]', 'QV:SEADATANET',
            'WindDirFrom [deg T]', 'QV:SEADATANET', 'QV:SEADATANET:SAMPLE']
MICRO_ITEMS = "SDN:H01::H0100001 (microplastic items)"
FRAGMENTS = "SDN:H01::H0100002 (fragments)"
FILAMENTS = "SDN:H01::H0100004 (filaments)"
NONPLASTIC = ("SDN:H01::H0100007 (non-plastic man-made micro-particles "
              "(e.g. glass, metal, tar))")
PE = "SDN:H05::H0500001 (polyethylene)"
S05_15 = "SDN:H03::H0300012 (0.5 to 1.5 millimetres)"
S15_5 = "SDN:H03::H0300013 (1.5 to 5 millimetres)"
S5_25 = "SDN:H03::H0300009 (5 to 25 millimetres)"
S03_5 = "SDN:H03::H0300010 (300 micrometres to 5 millimetres)"
CCBY = "Creative Commons Attribution 4.0 International(CB)"


def write_emodnet(root, pump_copy):
    """The ODV spreadsheet in its zip, and the truth of the samples kept.

    Samples (sampID): A plastic in two classes that tile 0.5-5 mm + a
    non-plastic line; B ambiguous fragments that caught something (dropped);
    C fragments of polyethylene + empty ambiguous filaments; D volume only, a
    measured zero; E a copy of an NCEI pump sample (dropped); F VOR-coded
    (dropped); G non-plastic only (dropped); H a date-only sample.
    """
    idx = {}
    for i, n in enumerate(EMOD_HDR):
        idx.setdefault(n, i)
    lines, truth = [], []

    def line(samp, when, lon, lat, gear, width, trk, vol, count, typ,
             size, poly="", weight="", vor="", wind="4", cruise="CRUISE-1",
             mind="0", maxd="0.5"):
        r = [""] * len(EMOD_HDR)
        r[idx["Cruise"]] = f"{samp}_ODV.txt"
        r[idx["Type"]] = "B"
        r[idx["yyyy-mm-ddThh:mm:ss.sss"]] = when
        r[idx["Longitude [degrees_east]"]] = f"{lon}"
        r[idx["Latitude [degrees_north]"]] = f"{lat}"
        r[idx["LOCAL_CDI_ID"]] = f"cdi_{samp}"
        r[idx["EDMO_code"]] = "120"
        r[idx["Data set name"]] = f"Floating micro-litter {samp}"
        r[idx["Instrument / gear type"]] = gear
        r[idx["Cruise name"]] = cruise
        r[idx["CSR platform name"]] = "RV Test"
        r[idx["Access restriction"]] = CCBY
        r[idx["MinDepBelowSeaSurf [m]"]] = mind
        r[idx["MaxDepBelowSeaSurf [m]"]] = maxd
        r[idx["mouth_width [cm]"]] = width
        r[idx["Mesh_size [micrometres]"]] = "333"
        r[idx["SampTrkLen [km]"]] = trk
        r[idx["VolWBodySamp [m3]"]] = vol
        r[idx["sampID"]] = samp
        r[idx["Microlitter_count_water"]] = count
        r[idx["Microlitter_count_water_VOR"]] = vor
        r[idx["*Microlitter_weight_water [g]"]] = weight
        r[idx["Microlitter_type_water"]] = typ
        r[idx["Microlitter_size_class_water"]] = size
        r[idx["Microlitter_polymer_type_water"]] = poly
        r[idx["WindSpd [m/s]"]] = wind
        lines.append("\t".join(r))

    def T(samp, y, m, d, hh, mi, lat, lon, gear, v, qc=0, cruise="CRUISE-1"):
        t = _t(y, m, d) - 43200 + hh * 3600 + mi * 60 if hh is not None \
            else _t(y, m, d)
        truth.append(("emodnet", (len(truth),), dict(
            t=t, lat=lat, lon=lon, qc=gear | qc,
            plat=f"emodnet|{cruise}|RV Test", v=v)))
    neu = "neuston nets(69)"
    area = 0.6 / 1000.0 * 2.0          # 60 cm x 2 km, in km2
    # A
    line("A", "2014-06-01T10:30:00", 5.0, 43.0, neu, "60", "2", "", "10",
         MICRO_ITEMS, S05_15, weight="0.01")
    line("A", "2014-06-01T10:30:00", 5.0, 43.0, neu, "60", "2", "", "20",
         MICRO_ITEMS, S15_5, weight="0.03")
    line("A", "2014-06-01T10:30:00", 5.0, 43.0, neu, "60", "2", "", "5",
         NONPLASTIC, S15_5, weight="0.5")
    T("A", 2014, 6, 1, 10, 30, 43.0, 5.0, 2,
      {"log_n_km2": L(30 / area), "log_g_km2": L(0.04 / area),
       "log_n_km2_micro": L(30 / area), "mesh_um": 333.0,
       "net_depth_m": 0.5, "sample_depth_m": 0.0, "wind_ms": 4.0})
    # B: ambiguous fragments that caught 3 -> no total
    line("B", "2014-06-02T10:30:00", 5.5, 43.5, neu, "60", "2", "", "3",
         FRAGMENTS, S03_5)
    line("B", "2014-06-02T10:30:00", 5.5, 43.5, neu, "60", "2", "", "4",
         MICRO_ITEMS, S03_5)
    # C: polyethylene fragments + empty ambiguous filaments
    line("C", "2014-06-03T08:00:00", 6.0, 44.0, "plankton nets(22)", "60",
         "2", "", "7", FRAGMENTS, S5_25, poly=PE)
    line("C", "2014-06-03T08:00:00", 6.0, 44.0, "plankton nets(22)", "60",
         "2", "", "0", FILAMENTS, S5_25)
    T("C", 2014, 6, 3, 8, 0, 44.0, 6.0, 7,
      {"log_n_km2": L(7 / area), "mesh_um": 333.0, "net_depth_m": 0.5,
       "sample_depth_m": 0.0, "wind_ms": 4.0})
    # D: volume only, a measured zero
    line("D", "2015-03-04T12:00:00", 7.0, 45.0, "unknown(999)", "", "",
         "10", "0", MICRO_ITEMS, S03_5, mind="0", maxd="0")
    T("D", 2015, 3, 4, 12, 0, 45.0, 7.0, 0,
      {"log_n_m3": 0.0, "mesh_um": 333.0, "sample_depth_m": 0.0,
       "wind_ms": 4.0})
    # E: a copy of an NCEI pump sample (same day, 0.005 deg away)
    y, m, d, la, lo = pump_copy
    line("E", f"{y:04d}-{m:02d}-{d:02d}T09:00:00", lo + 0.005, la - 0.005,
         neu, "60", "2", "", "9", MICRO_ITEMS, S03_5)
    # F: VOR-coded
    line("F", "2018-03-19T22:29:30", -55.8, -51.6,
         "continuous water samplers(31)", "", "", "0.037", "", MICRO_ITEMS,
         S03_5, vor="3.9975")
    # G: non-plastic only
    line("G", "2016-07-07T07:07:00", 8.0, 46.0, neu, "60", "2", "", "6",
         NONPLASTIC, S03_5)
    # H: date only -> noon, flagged
    line("H", "2016-05-03", 9.0, 47.0, neu, "60", "2", "", "2", MICRO_ITEMS,
         S03_5, cruise="")
    T("H", 2016, 5, 3, None, None, 47.0, 9.0, 2,
      {"log_n_km2": L(2 / area), "mesh_um": 333.0, "net_depth_m": 0.5,
       "sample_depth_m": 0.0, "wind_ms": 4.0}, qc=mp.QC_NO_TIME,
      cruise="Floating micro-litter H")
    text = ("//<Encoding>UTF-8</Encoding>\n//<Version>ODV Spreadsheet V4.8"
            "</Version>\n//\n" + "\t".join(EMOD_HDR) + "\n" +
            "\n".join(lines) + "\n")
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w") as z:
        z.writestr(mp.EMODNET_TXT, text)
        z.writestr("__MACOSX/._" + mp.EMODNET_TXT, b"decoy")
    _w(root, mp.EMODNET_ZIP, bio.getvalue())
    return truth


# ==================================================================== truth ==
def lebreton2018_truth(evs, hist):
    out = []
    for e in evs:
        t = _t(e["y"], e["mo"], e["d"], e["hour"])
        q = 0 if e["hour"] is not None else mp.QC_NO_TIME
        la = (e["lat0"] + e["lat1"]) / 2
        lo = (e["lon0"] + e["lon1"]) / 2
        expd = ("AERIAL expedition 2016" if e["type"] == "aerial"
                else "MEGA expedition 2015")
        base = {"sample_depth_m": 0.0, "wind_ms": 10.0 * KN}
        if e["depth_km"] is not None:
            base["net_depth_m"] = e["depth_km"] * 1000.0
        if e["type"] == "manta":
            mi, me = e["micro"], e["meso"]
            v = {"log_n_km2": L(mi["raw_n"] + me["raw_n"]),
                 "log_g_km2": L(mi["raw_g"] + me["raw_g"]),
                 "log_n_km2_wc": L(mi["mid_n"] + me["mid_n"]),
                 "log_n_km2_micro": L(mi["raw_n"]),
                 "log_n_km2_meso": L(me["raw_n"])}
            gear = 1
        elif e["type"] == "mega":
            c = e["macro"]
            v = {"log_n_km2": L(c["raw_n"]), "log_g_km2": L(c["raw_g"]),
                 "log_n_km2_wc": L(c["mid_n"]),
                 "log_n_km2_macro": L(c["raw_n"])}
            gear = 3
        else:
            c = e["mega"]
            v = {"log_n_km2": L(c["raw_n"]), "log_n_km2_wc": L(c["mid_n"]),
                 "log_n_km2_mega": L(c["raw_n"])}
            gear = 6
        v.update(base)
        out.append(("lebreton2018", (e["k"],), dict(
            t=t, lat=la, lon=lo, qc=gear | q,
            plat=f"lebreton2018|{expd}|{e['platform']}", v=v)))
    for i, h in enumerate(hist):
        if h["ref"] == "This study":
            continue                      # copies of the manta events
        gear = mp.HIST_GEAR[h["gear"] or None]
        dep = None if h["depth"] == "nan" else h["depth"]
        v = {"log_n_km2": L(h["n"]), "log_g_km2": L(h["g"]),
             "mesh_um": h["mesh"], "net_depth_m": dep,
             "sample_depth_m": 0.0 if gear in mp.AREAL_GEAR else None}
        ref = " ".join(h["ref"].split())
        out.append(("lebreton2018_hist", (i,), dict(
            t=_t_mid_month(h["y"], h["mo"]), lat=h["lat"], lon=h["lon"],
            qc=gear | mp.QC_NO_TIME | mp.QC_NO_DAY,
            plat=f"lebreton2018-historical|{ref}|{h['origin']}", v=v)))
    return out


def _finish(entries, d_lo, d_hi):
    """(src, subkey, row) -> the truth list in the store's tie order."""
    rank = {s: i for i, s in enumerate(mp.SOURCES)}
    lo = f10b.seconds_since_epoch(d_lo)
    hi = f10b.seconds_since_epoch(d_hi) + 86399
    entries = sorted(entries, key=lambda e: (e[2]["t"], rank[e[0]], e[1]))
    out = []
    for src, _k, r in entries:
        if not (lo <= r["t"] <= hi):
            continue
        v = [NAN] * len(mp.NAMES)
        for n, x in r["v"].items():
            v[mp.COL[n]] = NAN if x is None else float(x)
        if not any(np.isfinite(v[mp.COL[c]]) for c in mp.CONC):
            continue
        out.append({"t": int(r["t"]), "lat": float(r["lat"]),
                    "lon": float(r["lon"]),
                    "platform": f10b.platform_hash(r["plat"]), "v": v,
                    "qc": int(r["qc"]), "src": src})
    return out


def write_all(root, corrupt_swap=False):
    evs = events()
    hist, erik = historical(evs)
    objs = objects(evs)
    write_lebreton2018(root, evs, hist, objs, corrupt_swap=corrupt_swap)
    t24 = write_lebreton2024(root, evs)
    tn = write_ncei(root, evs, erik)
    te = write_emodnet(root, (2013, 4, 1, 55.0, -10.0))
    return evs, hist, objs, t24, tn, te


def make_smoke_sources(root, d_lo, d_hi, corrupt_swap=False):
    evs, hist, objs, t24, tn, te = write_all(root, corrupt_swap)
    entries = lebreton2018_truth(evs, hist) + t24 + tn + te
    # the Eriksen rows NCEI copies are dropped there, not here: the compiled
    # rows themselves are kept, and are in lebreton2018_truth already
    return _finish(entries, d_lo, d_hi)


def make_mobjects_smoke(root, d_lo, d_hi):
    evs, hist, objs, *_ = write_all(root)
    from family1.adapters import mobjects as mo
    lo = f10b.seconds_since_epoch(d_lo)
    hi = f10b.seconds_since_epoch(d_hi) + 86399
    out = []
    for o in objs:
        e = o["e"]
        t = _t(e["y"], e["mo"], e["d"], e["hour"])
        if not (lo <= t <= hi):
            continue
        out.append({"t": t, "lat": (o["lat0"] + o["lat1"]) / 2,
                    "lon": (o["lon0"] + o["lon1"]) / 2,
                    "platform": f10b.platform_hash(
                        f"lebreton2018|AERIAL expedition 2016|"
                        f"{e['platform']}|mosaic {lb.event_id(e['cid'])}"),
                    "v": [o["length"], o["width"],
                          NAN if o["area"] is None else o["area"],
                          o["kg"][0], o["kg"][1], o["kg"][2],
                          1.0 if o["lt50"] == "yes" else 0.0],
                    "qc": mo.CODE_OF[o["type"]] | mp.QC_NO_TIME})
    return out
